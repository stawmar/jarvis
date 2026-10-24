import json
import re
from pathlib import Path
from docx import Document

SOURCE_DIR = Path(
    "/srv/data/nextcloud/data/marek/files/"
    "Raporty/JARVIS_TEST/Po_OCR"
)

OUTPUT_FILE = Path("document_scan.json")


def read_docx(path):
    doc = Document(path)
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)


def extract_data(text):
    case_match = re.search(
        r"\b\d{4}-\d+\.Ds\.\d+\.\d{4}\b",
        text,
        re.IGNORECASE
    )

    internal_match = re.search(
        r"\b\d{4}/CHR/\d{2}\b",
        text,
        re.IGNORECASE
    )

    evidence = []

    for match in re.finditer(
        r"ślad/przedmiot\s+nr\s+(\d+)\s*[-–]\s*([^\n]+)",
        text,
        re.IGNORECASE
    ):
        evidence.append({
            "number": match.group(1),
            "description": match.group(2).strip()
        })

    return {
        "source": "BES" if "BIURO EKSPERTYZ SĄDOWYCH" in text.upper() else None,
        "caseNumber": case_match.group(0) if case_match else None,
        "internalNumber": internal_match.group(0) if internal_match else None,
        "evidence": evidence
    }


def main():
    results = []

    for path in sorted(SOURCE_DIR.glob("*.docx")):
        try:
            content = read_docx(path)
            data = extract_data(content)

            results.append({
                "filename": path.name,
                "status": "OK",
                "data": data,
                "ocrText": content,
                "approved": False
            })

            print(
                f"{path.name} | "
                f"sprawa: {data['caseNumber']} | "
                f"wew: {data['internalNumber']} | "
                f"dowody: {len(data['evidence'])}"
            )

        except Exception as exc:
            results.append({
                "filename": path.name,
                "status": "ERROR",
                "error": str(exc)
            })
            print(f"ERROR: {path.name}: {exc}")

    OUTPUT_FILE.write_text(
        json.dumps(results, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    print(f"\nPrzeanalizowano: {len(results)} dokumentów")
    print(f"Wynik zapisano: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()